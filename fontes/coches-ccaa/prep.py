#!/usr/bin/env python3
"""Turismos por 1.000 habitantes e peso dos «cero emisións» no parque e nas matriculacións → data/datos.json

Fontes:
  DGT, «Parque de vehículos. Tablas estadísticas 2025» (parque a 31/12/2025): táboas 64 (V.4.C.A,
    parque por comunidade e tipo) e 82 (P_TUR_MEDIO, parque de turismos por provincia e distintivo
    ambiental):
    https://www.dgt.es/menusecundario/dgt-en-cifras/dgt-en-cifras-resultados/dgt-en-cifras-detalle/Parque-de-vehiculos-Tablas-Estadisticas-2025/
    https://www.dgt.es/export/sites/web-DGT/.galleries/downloads/dgt-en-cifras/publicaciones/Parque-de-vehiculos-Tablas-Estadisticas/Parque-de-vehiculos-Tablas-estadisticas-2025.xlsx
  DGT, «Matriculaciones. Tablas estadísticas 2025»: táboas 42 (M_TUR_MEDIO, turismos matriculados
    por provincia e distintivo) e 41 (V_1_10_RENT_PROV, matriculacións de renting por provincia):
    https://www.dgt.es/export/sites/web-DGT/.galleries/downloads/dgt-en-cifras/publicaciones/matriculaciones-tablas-estadisticas/Matriculaciones-Tablas-estadisticas-2025.xlsx
  INE, Estatística Continua de Poboación, poboación residente a 1/1/2026 por comunidade
    (táboa 56940, series vía API Tempus):
    https://servicios.ine.es/wstempus/js/ES/DATOS_TABLA/56940?nult=1
    https://servicios.ine.es/wstempus/js/ES/DATOS_SERIE/<código>?nult=4

Precisa openpyxl (python3 -m pip install openpyxl).
"""
import json
import unicodedata
import warnings
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl

warnings.filterwarnings("ignore")
HERE = Path(__file__).resolve().parent
RAW = HERE / "data" / "raw"
RAW.mkdir(parents=True, exist_ok=True)
UA = {"User-Agent": "Mozilla/5.0"}
DGT = "https://www.dgt.es/export/sites/web-DGT/.galleries/downloads/dgt-en-cifras/publicaciones/"
FICH = {"parque-2025.xlsx": DGT + "Parque-de-vehiculos-Tablas-Estadisticas/Parque-de-vehiculos-Tablas-estadisticas-2025.xlsx",
        "matric-tablas-2025.xlsx": DGT + "matriculaciones-tablas-estadisticas/Matriculaciones-Tablas-estadisticas-2025.xlsx"}


def get(url):
    return urlopen(Request(url, headers=UA), timeout=300).read()


for f, u in FICH.items():
    if not (RAW / f).exists():
        (RAW / f).write_bytes(get(u))

sa = lambda s: unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower().replace(" ", "").replace("\n", "")
CCAA = [("AN", "Andalucía", ["almeria", "cadiz", "cordoba", "granada", "huelva", "jaen", "malaga", "sevilla"]),
        ("AR", "Aragón", ["huesca", "teruel", "zaragoza"]), ("AS", "Asturias", ["asturias"]),
        ("IB", "Illes Balears", ["balears(illes)"]), ("CN", "Canarias", ["palmas(las)", "santacruzdetenerife"]),
        ("CB", "Cantabria", ["cantabria"]),
        ("CL", "Castela e León", ["avila", "burgos", "leon", "palencia", "salamanca", "segovia", "soria", "valladolid", "zamora"]),
        ("CM", "Castela-A Mancha", ["albacete", "ciudadreal", "cuenca", "guadalajara", "toledo"]),
        ("CT", "Cataluña", ["barcelona", "girona", "lleida", "tarragona"]),
        ("VC", "C. Valenciana", ["alicante/alacant", "castellon/castello", "valencia/valencia"]),
        ("EX", "Estremadura", ["badajoz", "caceres"]), ("GA", "Galicia", ["coruna(a)", "lugo", "ourense", "pontevedra"]),
        ("MD", "Madrid", ["madrid"]), ("MC", "Murcia", ["murcia"]), ("NC", "Navarra", ["navarra"]),
        ("PV", "País Vasco", ["araba/alava", "gipuzkoa", "bizkaia"]), ("RI", "A Rioxa", ["rioja(la)"]),
        ("CE", "Ceuta", ["ceuta"]), ("ML", "Melilla", ["melilla"])]
PROV2CA = {p: c for c, _, ps in CCAA for p in ps}


def taboa(wb, folla, cab):
    """Filas por provincia {nome normalizado: {columna: valor}} da folla dada."""
    filas = list(wb[folla].iter_rows(values_only=True))
    i = next(k for k, r in enumerate(filas) if r[0] and str(r[0]).strip().startswith(cab))
    cols = [str(c).replace("\n", " ").strip() if c else "" for c in filas[i]]
    out = {}
    for r in filas[i + 1:]:
        if r[0] and isinstance(r[1], (int, float)):
            out[sa(str(r[0]))] = dict(zip(cols, r))
    return out


def por_ca(t, campos):
    acc = {c: {k: 0 for k in campos} for c, _, _ in CCAA}
    for p, r in t.items():
        if p in PROV2CA:
            for k in campos:
                acc[PROV2CA[p]][k] += r.get(k) or 0
    falta = set(PROV2CA) - set(t)
    assert not falta, falta
    return acc


DIST = ["CERO", "ECO", "C", "B", "Sin distintivo"]
wbp = openpyxl.load_workbook(RAW / "parque-2025.xlsx", read_only=True, data_only=True)
wbm = openpyxl.load_workbook(RAW / "matric-tablas-2025.xlsx", read_only=True, data_only=True)
parque = por_ca(taboa(wbp, "P_TUR_MEDIO", "Provincias"), DIST)
matric = por_ca(taboa(wbm, "M_TUR_MEDIO", "Provincias"), DIST)
rent_t = taboa(wbm, "V_1_10_RENT_PROV", "Provincias")
col_tur = next(k for k in next(iter(rent_t.values())) if k.lower().startswith("turismo"))
renting = por_ca(rent_t, [col_tur])

# Poboación a 1/1/2026 (INE, ECP)
NOMES_INE = {"Andalucía": "AN", "Aragón": "AR", "Asturias, Principado de": "AS", "Balears, Illes": "IB", "Canarias": "CN",
             "Cantabria": "CB", "Castilla y León": "CL", "Castilla - La Mancha": "CM", "Cataluña": "CT",
             "Comunitat Valenciana": "VC", "Extremadura": "EX", "Galicia": "GA", "Madrid, Comunidad de": "MD",
             "Murcia, Región de": "MC", "Navarra, Comunidad Foral de": "NC", "País Vasco": "PV", "Rioja, La": "RI",
             "Ceuta": "CE", "Melilla": "ML", "Total Nacional": "ES"}
ft = RAW / "ine_56940.json"
if not ft.exists():
    ft.write_bytes(get("https://servicios.ine.es/wstempus/js/ES/DATOS_TABLA/56940?nult=1"))
pob, data_pob = {}, None
for s in json.loads(ft.read_text(encoding="utf-8")):
    partes = [x.strip() for x in s["Nombre"].split(".")]
    if "Todas las edades" not in partes or ("Total" not in partes[:3] and partes[0] != "Total Nacional"):
        continue
    nome = next((p for p in partes if p in NOMES_INE), None)
    if not nome or (nome != "Total Nacional" and partes[0] != "Total"):
        continue
    if nome == "Total Nacional" and partes[2] != "Total":
        continue
    fs = RAW / f"ine_{s['COD']}.json"
    if not fs.exists():
        fs.write_bytes(get(f"https://servicios.ine.es/wstempus/js/ES/DATOS_SERIE/{s['COD']}?nult=4"))
    for x in json.loads(fs.read_text(encoding="utf-8"))["Data"]:
        d = datetime.fromtimestamp(x["Fecha"] / 1000, tz=timezone.utc).strftime("%Y-%m-%d")
        if d.endswith("-01-01") or d.endswith("12-31"):  # 1 de xaneiro (en hora local pode vir como 31/12 23:00 UTC)
            ano = d[:4] if d.endswith("-01-01") else str(int(d[:4]) + 1)
            if ano == "2026":
                pob[NOMES_INE[nome]] = x["Valor"]
                data_pob = "2026-01-01"
assert len(pob) == 20, pob

out = []
for c, n, _ in CCAA:
    p, m = parque[c], matric[c]
    tp, tm = sum(p.values()), sum(m.values())
    out.append({"c": c, "n": n, "pob": pob[c], "tur": tp, "tur1000": round(tp / pob[c] * 1000, 1),
                "p": {k: p[k] for k in DIST}, "m": {k: m[k] for k in DIST}, "mat": tm,
                "mat1000": round(tm / pob[c] * 1000, 2), "rent": renting[c][col_tur],
                "cero_p": round(p["CERO"] / tp * 100, 2), "cero_m": round(m["CERO"] / tm * 100, 2),
                "eco_m": round(m["ECO"] / tm * 100, 2), "sin_p": round(p["Sin distintivo"] / tp * 100, 2)})
tot = {k: sum(o["p"][k] for o in out) for k in DIST}
totm = {k: sum(o["m"][k] for o in out) for k in DIST}
es = {"pob": pob["ES"], "tur": sum(tot.values()), "mat": sum(totm.values()),
      "cero_p": round(tot["CERO"] / sum(tot.values()) * 100, 2), "cero_m": round(totm["CERO"] / sum(totm.values()) * 100, 2),
      "tur1000": round(sum(tot.values()) / pob["ES"] * 1000, 1), "rent": sum(o["rent"] for o in out)}
for o in sorted(out, key=lambda o: -o["tur1000"]):
    print(f"{o['n']:17s} {o['tur1000']:6.1f} tur/1000  CERO parque {o['cero_p']:5.2f}%  matric {o['cero_m']:5.2f}%  renting {o['rent']}")
print("España", es)
(HERE / "data" / "datos.json").write_text(json.dumps({"ano": 2025, "data_pob": data_pob, "es": es, "ccaa": out},
                                                     ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

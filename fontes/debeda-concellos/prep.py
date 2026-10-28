#!/usr/bin/env python3
"""Débeda viva de cada concello galego por habitante, 2025 e 2015 → data/datos.json

Fontes:
  · Ministerio de Facenda, «Deuda viva de las entidades locales», ficheiro de concellos a 31 de
    decembro (miles de euros). Páxina:
    https://www.hacienda.gob.es/es-ES/CDI/Paginas/SistemasFinanciacionDeuda/InformacionEELLs/DeudaViva.aspx
      2025: https://www.hacienda.gob.es/cdi/sist%20financiacion%20y%20deuda/informacioneells/2025/deuda-viva-ayuntamientos-202512.xlsx
      2015: https://www.hacienda.gob.es/sgfal/dgcfel/deudaviva/deuda%20viva%20ayuntamientos%202015__ovel_20160506.xls
    É a débeda financeira do concello (préstamos, emisións, operacións con entidades de crédito)
    segundo a Central de Información de Riscos das Entidades Locais (CIR-Local); non inclúe a débeda
    comercial con provedores nin a das deputacións.
  · INE, Cifras oficiais de poboación (Padrón) por concello, táboas 2868 (A Coruña), 2880 (Lugo),
    2885 (Ourense) e 2890 (Pontevedra): https://www.ine.es/jaxiT3/files/t/es/csv_bdsc/<táboa>.csv
    Poboación a 1/1/2025 para a débeda de 2025 e a 1/1/2016 para a de 2015.
Fusións: Oza-Cesuras (15902, 2013) e Cerdedo-Cotobade (36902 = 36011 + 36012, 2016); en 2015 súmanse.
Precisa openpyxl e xlrd (pip install openpyxl xlrd). Ficheiros gardados en data/raw/.
"""
import csv
import json
from pathlib import Path
from urllib.request import Request, urlopen

import openpyxl
import xlrd

HERE = Path(__file__).resolve().parent
RAW = HERE / "data" / "raw"
RAW.mkdir(parents=True, exist_ok=True)
UA = {"User-Agent": "Mozilla/5.0"}
F = {
    "dv-ayto-202512.xlsx": "https://www.hacienda.gob.es/cdi/sist%20financiacion%20y%20deuda/informacioneells/2025/deuda-viva-ayuntamientos-202512.xlsx",
    "dv-ayto-2015.xls": "https://www.hacienda.gob.es/sgfal/dgcfel/deudaviva/deuda%20viva%20ayuntamientos%202015__ovel_20160506.xls",
    **{f"ine_{t}.csv": f"https://www.ine.es/jaxiT3/files/t/es/csv_bdsc/{t}.csv" for t in (2868, 2880, 2885, 2890)},
}
for n, u in F.items():
    if not (RAW / n).exists():
        (RAW / n).write_bytes(urlopen(Request(u, headers=UA), timeout=300).read())
FUS = {"36011": "36902", "36012": "36902", "15026": "15902", "15063": "15902"}
PROV = {"15": "A Coruña", "27": "Lugo", "32": "Ourense", "36": "Pontevedra"}


def filas_debeda(n):
    if n.endswith(".xlsx"):
        ws = openpyxl.load_workbook(RAW / n, read_only=True, data_only=True)["Datos"]
        rows = ws.iter_rows(values_only=True)
    else:
        s = xlrd.open_workbook(RAW / n).sheet_by_name("Datos")
        rows = (s.row_values(i) for i in range(s.nrows))
    out = {}
    for r in rows:
        if r and r[2] == "GALICIA" and str(r[3]) in PROV:
            k = f"{r[3]}{r[5]}"
            k = FUS.get(k, k)
            out[k] = out.get(k, 0) + float(r[7] or 0) * 1000
    return out


d25, d15 = filas_debeda("dv-ayto-202512.xlsx"), filas_debeda("dv-ayto-2015.xls")


def bonito(n):
    if ", " in n:
        a, b = n.rsplit(", ", 1)
        if b in ("A", "O", "As", "Os"):
            return f"{b} {a}"
    return n


pob, nomes = {"2025": {}, "2016": {}}, {}
for t in (2868, 2880, 2885, 2890):
    for r in csv.reader((RAW / f"ine_{t}.csv").open(encoding="utf-8-sig"), delimiter=";"):
        if len(r) < 4 or r[1] != "Total" or r[2] not in pob or not r[0][:5].isdigit() or not r[3].strip('"'):
            continue
        k = FUS.get(r[0][:5], r[0][:5])
        pob[r[2]][k] = pob[r[2]].get(k, 0) + int(r[3].replace(".", ""))
        if r[0][:5] == k:
            nomes[k] = bonito(r[0][6:])
conc = []
for k in sorted(pob["2025"]):
    conc.append({"id": k, "n": nomes[k], "pr": PROV[k[:2]], "p": pob["2025"][k], "d": round(d25.get(k, 0)),
                 "p15": pob["2016"][k], "d15": round(d15.get(k, 0))})
assert len(conc) == 313, len(conc)
falta = [k for k in pob["2025"] if k not in d25]
print("sen fila de débeda en 2025 (cóntanse como 0):", falta)
pc = lambda c: c["d"] / c["p"]
cero = [c for c in conc if c["d"] == 0]
print(len(cero), "concellos con débeda cero en 2025;", sum(1 for c in conc if c["d15"] == 0), "en 2015")
tot = sum(c["d"] for c in conc)
print(f"total {tot / 1e6:.1f} M€, {tot / sum(c['p'] for c in conc):.0f} €/hab; 2015: {sum(c['d15'] for c in conc) / 1e6:.1f} M€")
for c in sorted(conc, key=pc, reverse=True)[:10]:
    print(f"  {c['n']:24} {pc(c):6.0f} €/hab  {c['d'] / 1e6:6.1f} M€   (2015: {c['d15'] / c['p15']:.0f})")
# Agregado de España (concellos) para contexto
s = openpyxl.load_workbook(RAW / "dv-ayto-202512.xlsx", read_only=True, data_only=True)["Datos"]
esp = sum(float(r[7] or 0) for r in s.iter_rows(values_only=True) if r and isinstance(r[7], (int, float)) and r[0] == "2025") * 1000
print(f"España: {esp / 1e9:.1f} mil M€")
out = {"ano": 2025, "ano0": 2015, "conc": conc, "esp": round(esp)}
(HERE / "data" / "datos.json").write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print("ok: data/datos.json")
